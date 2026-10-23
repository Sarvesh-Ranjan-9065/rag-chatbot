# utils/extractor.py
import os, io, logging
from azure.ai.documentintelligence import DocumentIntelligenceClient
from azure.core.credentials import AzureKeyCredential

# File types handled by Azure Document Intelligence (OCR-capable)
DI_EXTENSIONS = {'.pdf', '.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.heif'}

def get_di_client():
    return DocumentIntelligenceClient(
        endpoint=os.environ["AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT"],
        credential=AzureKeyCredential(os.environ["AZURE_DOCUMENT_INTELLIGENCE_KEY"])
    )

def extract_text(blob_data: bytes, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    logging.info(f"Extracting text from: {filename} (type: {ext})")

    try:
        # ── PDFs and Images → Azure Document Intelligence ──────────
        if ext in DI_EXTENSIONS:
            return _extract_with_di(blob_data)

        # ── Word Documents ──────────────────────────────────────────
        elif ext == '.docx':
            return _extract_docx(blob_data)

        # ── PowerPoint ─────────────────────────────────────────────
        elif ext == '.pptx':
            return _extract_pptx(blob_data)

        # ── Excel ───────────────────────────────────────────────────
        elif ext in ('.xlsx', '.xls'):
            return _extract_xlsx(blob_data)

        # ── Plain text formats ──────────────────────────────────────
        elif ext in ('.txt', '.md', '.csv', '.json', '.xml', '.html', '.py', '.js'):
            return blob_data.decode('utf-8', errors='ignore')

        # ── Unknown: attempt UTF-8 decode ───────────────────────────
        else:
            logging.warning(f"Unknown extension {ext}, attempting UTF-8 decode")
            return blob_data.decode('utf-8', errors='ignore')

    except Exception as e:
        logging.error(f"Extraction failed for {filename}: {e}")
        return ""

def _extract_with_di(blob_data: bytes) -> str:
    client = get_di_client()
    poller = client.begin_analyze_document(
        model_id="prebuilt-read",
        body=blob_data,
        content_type="application/octet-stream"
    )
    result = poller.result()
    return result.content or ""

def _extract_docx(blob_data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(blob_data))
    return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

def _extract_pptx(blob_data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(blob_data))
    text_parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text)
    return "\n".join(text_parts)

def _extract_xlsx(blob_data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob_data), data_only=True)
    text_parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                text_parts.append(row_text)
    return "\n".join(text_parts)